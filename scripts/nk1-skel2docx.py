#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""NK1 骨架 → Jinja docx 模板 → 假值渲染（docxtpl）。

用法（須在 .venv：`. .venv/bin/activate`，相依見 requirements-docx.txt）：
  nk1-skel2docx.py SKEL.yaml --template-out T.docx --out R.docx

輸入：nk1-skel.py 產的骨架 YAML（本檔自帶小型讀取器，只支援 nk1-skel.py 的輸出形狀；本線未裝 PyYAML）。
  用 folded[]（repeat 已摺）排節；repeat 區塊的節層級與表格取自首次出現的 sections[]。
模板：
  - 檔頭（頁首＋首段）寫死「合成範例・非模板」
  - 節 → Heading N 段落
  - 表 → 表頭列（若 header）＋ `{%tr for row in t_k %}` 列迴圈；每列末格含 `{field:…}` 佔位符
  - repeat → `{%p for rep in repeat_k %}` … `{%p endfor %}`
      有 head（v1.1）：每次先出一個 head 標題 `{{ rep.label }}`，label＝該次對應的 head；block 節名不加後綴
      無 head：區塊內節名後綴 `{{ rep.label }}`（RUN-n）
假值：全為合成字串（VAL-…／RUN-n）；`{field:…}` 單大括號非 Jinja 語法，**須原樣保留**。
確定性（M1 D 段）：存檔後重寫 zip——所有 entry 的 mtime 固定 1980-01-01、entry 依名稱排序、屬性歸零；
  同 YAML＋同假值 → 模板與渲染檔 sha256 皆相同。
🔴 真骨架只准在 ~/nk1-data/ 下渲染，產物不進 repo。
"""
import argparse, json, os, re, sys, zipfile

HEADER = "合成範例・非模板"


# ── 骨架 YAML 讀取（nk1-skel.py 輸出子集）──────────────────────────
def _flow(v):
    v = v.strip()
    if not v:
        return None
    if v[0] in "[{":
        j = re.sub(r'([{,]\s*)([A-Za-z_][A-Za-z0-9_]*)\s*:', r'\1"\2":', v)
        return json.loads(j)
    if v[0] == '"':
        return json.loads(v)
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    return {"true": True, "false": False}.get(v, v)


def load_skel(path):
    d, key = {}, None
    for ln, line in enumerate(open(path, encoding="utf8"), 1):
        line = line.rstrip("\n")
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        m = re.match(r"^\s+-\s+(.*)$", line)
        if m and key:
            d[key].append(_flow(m.group(1)))
            continue
        m = re.match(r"^([A-Za-z_][A-Za-z0-9_]*):\s*(.*)$", line)
        if not m:
            raise ValueError(f"骨架 YAML 第 {ln} 行無法解析")          # 不回顯內容（真骨架可能含節名）
        k, v = m.group(1), m.group(2)
        if v.strip():
            d[k], key = _flow(v), None
        else:
            d[k], key = [], k
    for need in ("sections", "folded"):
        if need not in d:
            raise ValueError(f"骨架 YAML 缺 {need}")
    d.setdefault("repeats", [])
    d.setdefault("tables_before_first_section", [])
    return d


# ── 模板建構（python-docx）──────────────────────────────────────
def build_template(skel, path):
    from docx import Document
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = HEADER
    doc.add_paragraph(HEADER)
    ctx_shapes = {}                                   # 表變數名 → (資料列數, 欄數)
    tcount = [0]

    def add_table(t):
        tcount[0] += 1
        k = f"t_{tcount[0]}"
        rows, cols, header = int(t["rows"]), int(t["cols"]), bool(t["header"])
        nrows = (1 if header else 0) + 3              # 表頭？＋ for 列＋資料列＋ endfor 列
        tbl = doc.add_table(rows=nrows, cols=max(cols, 1))
        r = 0
        if header:
            for c in range(cols):
                tbl.cell(0, c).text = f"COL-{c + 1}"
            r = 1
        tbl.cell(r, 0).text = f"{{%tr for row in {k} %}}"
        for c in range(cols):
            tbl.cell(r + 1, c).text = f"{{{{ row.c{c + 1} }}}}"
        tbl.cell(r + 1, cols - 1).text += " {field:" + f"{k}_note" + "}"
        tbl.cell(r + 2, 0).text = "{%tr endfor %}"
        ctx_shapes[k] = (rows - (1 if header else 0), cols)

    def heading(name, level, suffix=""):
        doc.add_heading(name + suffix, level=min(max(int(level), 1), 9))

    for t in skel["tables_before_first_section"]:
        add_table(t)
    by_order = {s["order"]: s for s in skel["sections"]}
    reps = skel["repeats"]
    seq = skel["sections"]
    idx = 0                                           # sections 游標（folded 依序展開）
    for item in skel["folded"]:
        if isinstance(item, dict):
            k = item["repeat"]
            rp = reps[k - 1]
            start = rp["starts"][0]
            heads = rp.get("head")
            doc.add_paragraph(f"{{%p for rep in repeat_{k} %}}")
            if heads:
                h = by_order[start - 1]
                heading("{{ rep.label }}", h["level"])
                for t in h["tables"]:
                    add_table(t)
            for j in range(len(rp["block"])):
                s = by_order[start + j]
                heading(s["name"], s["level"], "" if heads else " {{ rep.label }}")
                for t in s["tables"]:
                    add_table(t)
            doc.add_paragraph("{%p endfor %}")
            ctx_shapes[f"repeat_{k}"] = (rp["count"], 0, heads)
            # 跳過所有出現位置（含 head）
            skip = {st + j for st in rp["starts"] for j in range(-1 if heads else 0, len(rp["block"]))}
            while idx < len(seq) and seq[idx]["order"] in skip:
                idx += 1
        else:
            while idx < len(seq) and seq[idx]["name"] != item:
                idx += 1
            s = seq[idx] if idx < len(seq) else {"name": item, "level": 1, "tables": []}
            heading(s["name"], s["level"])
            doc.add_paragraph("正文：{field:" + f"sec_{s.get('order', 0)}_body" + "}")
            for t in s["tables"]:
                add_table(t)
            idx += 1
    doc.save(path)
    return ctx_shapes


FIXED_DT = (1980, 1, 1, 0, 0, 0)


def normalize_zip(path):
    """zip 確定化：entry 依名稱排序、mtime＝1980-01-01、create_system／external_attr 固定、一律 deflate。"""
    with zipfile.ZipFile(path) as z:
        items = sorted((i.filename, z.read(i.filename)) for i in z.infolist())
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w") as z:
        for name, data in items:
            zi = zipfile.ZipInfo(name, date_time=FIXED_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 0
            zi.external_attr = 0
            z.writestr(zi, data)
    os.replace(tmp, path)


def fake_context(shapes):
    ctx = {}
    for k, shp in shapes.items():
        n, cols = shp[0], shp[1]
        if k.startswith("repeat_"):
            heads = shp[2] if len(shp) > 2 else None
            ctx[k] = [{"label": heads[i] if heads else f"RUN-{i + 1}"} for i in range(n)]
        else:
            ctx[k] = [{f"c{c + 1}": f"VAL-{k}-{r + 1}-{c + 1}" for c in range(cols)} for r in range(max(n, 0))]
    return ctx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("skel")
    ap.add_argument("--template-out", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    try:
        from docxtpl import DocxTemplate
    except ImportError:
        print("🔴 缺 docxtpl：請先 `. .venv/bin/activate`（requirements-docx.txt）", file=sys.stderr)
        sys.exit(2)
    try:
        skel = load_skel(a.skel)
        shapes = build_template(skel, a.template_out)
        tpl = DocxTemplate(a.template_out)
        tpl.render(fake_context(shapes))
        tpl.save(a.out)
        normalize_zip(a.template_out)
        normalize_zip(a.out)
    except Exception as e:
        print(f"🔴 渲染失敗：{type(e).__name__}: {str(e)[:120]}", file=sys.stderr)
        sys.exit(1)
    print(f"✅ 模板 {a.template_out}　渲染 {a.out}　表變數 {sum(1 for k in shapes if k.startswith('t_'))}　repeat {sum(1 for k in shapes if k.startswith('repeat_'))}")


if __name__ == "__main__":
    main()
