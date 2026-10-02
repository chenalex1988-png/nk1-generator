#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""NK1 重渲染（工單【M5.4】C；品保 2026-09-30 函 2-3）— 不啟模型；以既有生成品重出 docx，引用搬到審閱層。

流程（每個非固定節）：
  1. 先過 M5.4 B 結果陳述閘（全文內容判）：命中 → 攔下（佔位），不改寫。
  2. 正文含條文譯文（與條號表中譯連續 ≥15 字相同）→ 攔下（刪句＝改內容，違反改寫三關，不做）。
  3. 正文條號引用片段（「PIC/S GMP Annex 15（PE 009-18）X.Y…」及其「依／符合…之規定」外殼）→ 自正文移除，
     條件：片段內每個條號都在條號表 v1；記 normalizations{kind:"citation_moved", original, clauses, table_check}。
     任一條號不在條號表 → 不移除、攔下。
  4. 外部標準允收數值旁註：intake 允收項目帶 `source`（如「PE 009-18 §3.x」）者，於該數值旁註「（來源：…）」；無 source 者不註。
docx：正文零條號；新增「參考資料」節（法規名＋版次，取自本份所用 refs）置於「修改歷程」前；固定文節照舊。
另出 review_map.md（節 ↔ 條號審閱對照；條號→章節不再作正文攔截）與內容包（refs 全留＋normalizations）。
"""
import argparse, hashlib, importlib.util, json, os, re, sys, zipfile

VERSION = "1.0.0"
HERE = os.path.dirname(os.path.abspath(__file__))
_sp = importlib.util.spec_from_file_location("nk1_gen", os.path.join(HERE, "nk1-gen.py"))
g = importlib.util.module_from_spec(_sp)
_sp.loader.exec_module(g)

CL = r"\d{1,2}\.\d{1,2}"
LAWV = r"(?:PIC/S\s*GMP\s*)?Annex\s*15\s*[（(]\s*PE\s*009-18\s*[）)]\s*"
CIT = LAWV + CL + r"(?:(?:\s*[、,，及與和]\s*|\s+)(?:" + LAWV + r")?" + CL + r")*"
RE_WRAPPED = re.compile(r"[，,]?\s*(?:並且|並|且)?\s*(?:符合|依據|依照|參照|參考|根據|遵循|依)\s*" + CIT + r"\s*(?:中)?(?:之|的)?(?:相關)?(?:規定|要求|條文|原則|精神)?")
RE_BARE = re.compile(CIT)
BLOCKED_RS = "〔本節含結果陳述，已依品保 2026-09-30 函 2-1 攔下，待重新生成〕"
BLOCKED_TR = "〔本節含法規條文譯文，已依品保 2026-09-30 函 2-3 攔下，待重新生成〕"
BLOCKED_NT = "〔本節引用之條號不在條號表，未移除、已攔下〕"
BLOCKED_EM = "〔本節條號引用嵌在句中，移除會改變語句，未移除、已攔下，待重新生成〕"
LEFT_OK, RIGHT_OK = set("，,。；、（(：:"), set("，,。；）)：:")


def clean(t):
    t = re.sub(r"[，,]\s*([。；])", r"\1", t)
    t = re.sub(r"^[，,、\s]+", "", t)
    t = re.sub(r"([。；])[，,、\s]+", r"\1", t)
    t = re.sub(r"\s{2,}", " ", t)
    return t.strip()


def safe_boundary(text, m):
    """片段兩側須為句界或標點（或片段自帶前導逗號），否則移除會留下斷句——不安全。"""
    left = m.start() == 0 or text[m.start() - 1] in LEFT_OK or m.group(0)[:1] in "，,"
    right = m.end() == len(text) or text[m.end()] in RIGHT_OK
    return left and right


def strip_citations(sid, body, table):
    """回傳 (新正文, normalizations, 狀態)。狀態：ok／not_in_table／embedded（後兩者原文不改）。
    依序處理：帶外殼片段（依／符合…之規定）→ 裸片段；每一輪先全數檢查（邊界、條號在表）再移除。"""
    norms, out = [], body
    for rx in (RE_WRAPPED, RE_BARE):
        ms = list(rx.finditer(out))
        for m in ms:
            frag = m.group(0)
            cls = re.findall(CL, re.sub(LAWV, "", frag))
            rec = {"section_id": sid, "kind": "citation_moved", "original": frag, "clauses": cls}
            miss = [c for c in cls if c not in table]
            if miss:
                return body, [dict(rec, table_check=f"missing:[{','.join(miss)}]", action="block")], "not_in_table"
            if not safe_boundary(out, m):
                return body, [dict(rec, table_check="embedded", action="block")], "embedded"
            norms.append(dict(rec, table_check="all_in_table"))
        out = rx.sub("", out)
    return clean(out), norms, "ok"


def copies_translation(body, refs_v1, n=15):
    b = re.sub(r"\s", "", body)
    for r in refs_v1:
        t = re.sub(r"\s", "", r.get("translation_zh") or "")
        if any(t[i:i + n] in b for i in range(0, max(0, len(t) - n + 1))):
            return r["clause"]
    return None


def annotate_external(body, intake):
    n = 0
    for k in ("accept_iq", "accept_oq", "accept_pq"):
        for it in intake.get(k) or []:
            src = it.get("source")
            if not src:
                continue
            for v in g.cs.CATS["數值"].findall(it.get("criterion", "")):
                pass
            for m in g.cs.CATS["數值"].finditer(it.get("criterion", "")):
                val = m.group(0)
                if val in body and f"{val}（來源：{src}）" not in body:
                    body = body.replace(val, f"{val}（來源：{src}）", 1)
                    n += 1
    return body, n


def main():
    ap = argparse.ArgumentParser()
    for k in ("gen-dir", "skel", "refs", "intake", "intake-spec", "section-uses", "out-dir", "inbox-name"):
        ap.add_argument("--" + k, required=True)
    ap.add_argument("--run", default="run1")
    ap.add_argument("--inbox", default=os.path.join(os.path.expanduser("~"), "ak-inbox", "nk1-out"))
    a = ap.parse_args()
    skel = g.s2d.load_skel(a.skel)
    roles = g.section_roles(skel)
    by = {s["order"]: s for s in skel["sections"]}
    refs_v1 = g.load_refs(a.refs)
    table = {r["clause"] for r in refs_v1}
    intake = json.load(open(a.intake, encoding="utf8"))
    base = g.policy_base(intake, g.fy.load(a.intake_spec), g.fy.load(a.section_uses), refs_v1)
    src_dir = os.path.join(os.path.expanduser(a.gen_dir), a.run)
    secs, norms, stats, review = [], [], {"blocked_result": 0, "blocked_translation": 0, "blocked_not_in_table": 0, "blocked_embedded": 0,
                                           "citations_moved": 0, "external_notes": 0, "fixed": 0, "model": 0}, []
    used_laws = set()
    for l in open(os.path.join(src_dir, "sections.jsonl"), encoding="utf8"):
        x = json.loads(l)
        o, sid = x["order"], x["sid"]
        try:
            raw_refs = json.loads(x.get("raw") or "{}").get("refs") or []
        except Exception:
            raw_refs = []
        for r in raw_refs:
            used_laws.add((g.law_base(r.get("law", "")), r.get("version") or g.REF_VERSION))
        if x.get("fixed"):
            stats["fixed"] += 1
            secs.append({"order": o, "sid": sid, "body": x["body"], "fixed": True, "refs": [], "status": "fixed"})
            continue
        stats["model"] += 1
        body, status, nrm = x["body"], "ok", []
        ctx = g.policy_ctx(base, by[o], roles.get(o))
        rs = g.cs.result_statement_hits(body) if ctx.get("doc_kind") == "plan" else []
        trc = copies_translation(body, refs_v1)
        if rs:
            body, status = BLOCKED_RS, "blocked_result"
        elif trc:
            body, status = BLOCKED_TR, "blocked_translation"
        else:
            nb, nrm, st = strip_citations(sid, body, table)
            if st != "ok":
                body, status = (BLOCKED_NT, "blocked_not_in_table") if st == "not_in_table" else (BLOCKED_EM, "blocked_embedded")
            else:
                body, k = annotate_external(nb, intake)
                stats["external_notes"] += k
                stats["citations_moved"] += len(nrm)
        if status == "blocked_embedded":
            nrm = []
        if status != "ok":
            stats[status] += 1
        norms += nrm
        cited = sorted({c for n in nrm for c in n["clauses"]} | {str(r.get("clause")) for r in raw_refs}, key=lambda c: [int(p) for p in c.split(".")])
        review.append((sid, by[o]["name"], cited, sorted(ctx["clauses"], key=lambda c: [int(p) for p in c.split(".")])))
        secs.append({"order": o, "sid": sid, "body": body, "refs": raw_refs, "status": status})
    run_id = os.path.basename(os.path.normpath(os.path.expanduser(a.gen_dir)))
    out = os.path.join(os.path.expanduser(a.out_dir), run_id + "_v0.10")
    os.makedirs(out, exist_ok=True)
    tpl, docx = os.path.join(out, "template.docx"), os.path.join(out, "T027_draft.docx")
    g.render(skel, [{"order": s["order"], "body": s["body"]} for s in secs], tpl, docx)
    laws = sorted(f"{lw}（{v}）" for lw, v in used_laws if lw)
    add_reference_section(docx, laws)
    g.s2d.normalize_zip(docx)
    check = verify_docx(docx)
    with open(os.path.join(out, "review_map.md"), "w", encoding="utf8") as f:
        f.write("# 審閱對照：節 ↔ 條號（M5.4 C；條號→章節不作正文攔截，供審閱）\n\n"
                "| 節 | 節名 | 本節引用條號（refs＋自正文移出） | 條號表對應本節之條號 | 引用皆在對應內 |\n|---|---|---|---|---|\n")
        for sid, nm, cited, allowed in review:
            f.write(f"| {sid} | {nm} | {'、'.join(cited) or '—'} | {'、'.join(allowed) or '—'} | {'✅' if set(cited) <= set(allowed) else '⚠️'} |\n")
    pack = {"origin": "ai_draft", "disclosure": "本草稿由本地模型依合成資料逐節生成，未經人員修訂、審核與簽核，不得作為受控文件使用。",
            "validator_version": g.VALIDATOR, "rerender": f"nk1-rerender.py v{VERSION}", "source_gen": f"{run_id}/{a.run}",
            "docx_sha256": g.fsha(docx), "references_section": laws, "normalizations": norms,
            "sections": [{"sid": s["sid"], "status": s["status"], "body_sha256": g.sha(s["body"].encode()), "refs": s["refs"]} for s in secs]}
    pj = json.dumps(pack, ensure_ascii=False, sort_keys=True, indent=1).encode()
    open(os.path.join(out, "content_pack.json"), "wb").write(pj)
    os.makedirs(a.inbox, exist_ok=True)
    ib = os.path.join(a.inbox, a.inbox_name)
    open(ib, "wb").write(open(docx, "rb").read())
    ibp = re.sub(r"\.docx$", "", ib)
    open(ibp + ".content_pack.json", "wb").write(pj)
    open(ibp + ".review_map.md", "w", encoding="utf8").write(open(os.path.join(out, "review_map.md"), encoding="utf8").read())
    stats.update(check, docx_sha8=g.fsha(ib)[:8], pack_sha8=g.sha(pj)[:8], out=out, laws=laws)
    json.dump(stats, open(os.path.join(out, "stats.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)
    print(json.dumps(stats, ensure_ascii=False))


def add_reference_section(path, laws):
    """「參考資料」節：法規名＋版次；置於「修改歷程」標題前（找不到則置文末）。"""
    from docx import Document
    d = Document(path)
    anchor = next((p for p in d.paragraphs if (p.style.name or "").lower().startswith("heading") and p.text.startswith("修改歷程")), None)
    if anchor is not None:
        h = anchor.insert_paragraph_before("參考資料", style=anchor.style)
        for lw in laws:
            anchor.insert_paragraph_before(lw)
    else:
        d.add_heading("參考資料", level=1)
        for lw in laws:
            d.add_paragraph(lw)
    d.save(path)


def verify_docx(path):
    """正文（參考資料節除外）零條號、零法規名＋版次；參考資料節存在且有內容。"""
    from docx import Document
    d = Document(path)
    in_ref, body_cits, ref_lines, heads = False, 0, 0, 0
    for p in d.paragraphs:
        is_head = (p.style.name or "").lower().startswith("heading")
        if is_head:
            heads += 1
            in_ref = p.text.strip() == "參考資料"
            continue
        if in_ref:
            ref_lines += 1 if p.text.strip() else 0
        elif re.search(r"Annex\s*15|PE\s*009-18|§\s*\d", p.text):
            body_cits += 1
    return {"body_citation_paragraphs": body_cits, "reference_section_lines": ref_lines, "headings": heads}


if __name__ == "__main__":
    main()
